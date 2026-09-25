"""Create journal-neutral Word manuscript from the audited Markdown draft."""
from pathlib import Path
import re
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
src=ROOT/'manuscript/Manuscript_draft.md'
out=ROOT/'manuscript/Manuscript_draft.docx'
doc=Document();sec=doc.sections[0];sec.top_margin=Inches(.85);sec.bottom_margin=Inches(.85);sec.left_margin=Inches(.95);sec.right_margin=Inches(.95)
normal=doc.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(10.5);normal.font.color.rgb=RGBColor(0,0,0)
normal.paragraph_format.line_spacing=1.15;normal.paragraph_format.space_after=Pt(5)
for sty,size in [('Title',16),('Heading 1',12),('Heading 2',11)]:
    s=doc.styles[sty];s.font.name='Times New Roman';s.font.size=Pt(size);s.font.bold=True;s.font.color.rgb=RGBColor(0,0,0)
    s.paragraph_format.space_before=Pt(11 if sty!='Title' else 0);s.paragraph_format.space_after=Pt(5);s.paragraph_format.keep_with_next=True
    if sty=='Title':
        ppr=s.element.pPr
        if ppr is not None:
            border=ppr.find(qn('w:pBdr'))
            if border is not None:ppr.remove(border)
def add_inline(p,text):
    bits=re.split(r'(\*\*.*?\*\*)',text)
    for bit in bits:
        if bit.startswith('**') and bit.endswith('**'):
            run=p.add_run(bit[2:-2]);run.bold=True
        else:p.add_run(bit)
for line in src.read_text(encoding='utf-8').splitlines():
    line=line.rstrip()
    if not line:continue
    if line.startswith('# '):
        p=doc.add_paragraph(style='Title');p.add_run(line[2:]);p.alignment=WD_ALIGN_PARAGRAPH.LEFT
    elif line.startswith('## '):doc.add_paragraph(line[3:],style='Heading 1')
    elif line.startswith('### '):doc.add_paragraph(line[4:],style='Heading 2')
    else:
        p=doc.add_paragraph();add_inline(p,line.replace('  ',' '))
        if line.startswith(('1. ','2. ','3. ','4. ','5. ','6. ','7. ','8. ','9. ','10. ')):
            p.paragraph_format.space_after=Pt(2);p.style=normal
doc.add_page_break()
for num,name in [(1,'Figure_1_bulk_composition_and_survival.png'),(2,'Figure_2_spatial_coenrichment.png'),(3,'Figure_3_virtual_immune_validation.png')]:
    p=doc.add_paragraph(style='Heading 1');p.add_run(f'Figure {num}')
    pic=doc.add_paragraph();pic.alignment=WD_ALIGN_PARAGRAPH.CENTER
    pic.add_run().add_picture(str(ROOT/'manuscript/figures'/name),width=Inches(6.35))
    if num==1:doc.add_paragraph('Bulk cohort correlations and cohort-specific survival estimates. Exact values are provided in the results tables.')
    elif num==2:doc.add_paragraph('Patient-level median spatial correlations before and after adjustment. These represent RNA co-enrichment, not metabolite flux.')
    else:doc.add_paragraph('MCP-counter estimates support a shared immune-stromal association; conditional results vary by cohort.')
    if num<3:doc.add_page_break()
doc.save(out)
print(out)
